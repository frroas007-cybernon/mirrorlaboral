"""
Generador de documentos del Mirror 2 -- v3.
Copiado tal cual desde generar_documentos_v3.py (documento aprobado del
proyecto KARIN, ya revisado linea por linea) -- no se resume ni se reescribe
el contenido de negocio. Unico cambio respecto al original: el import de
`scoring` apunta al paquete app.mirror2_engine para poder importarse desde
la app web (misma logica, distinta ubicacion de archivo).
"""
from docx import Document
from docx.shared import Pt, RGBColor
from app.mirror2_engine.scoring import compute_scores, top_and_secondary, level_label, DIMENSION_NAMES

AZUL = RGBColor(0x17, 0x3F, 0x46)
CARBON = RGBColor(0x30, 0x34, 0x36)
GRIS = RGBColor(0x6B, 0x62, 0x59)

LECTURA_DIMENSION = {
    "trato": "Señales de trato humillante o descalificador, especialmente frente a terceros.",
    "presion": "Exigencias percibidas como desproporcionadas y sostenidas.",
    "aislamiento": "Percepción de exclusión de espacios o conversaciones antes habituales.",
    "amenazas": "Componente de miedo a consecuencias o vigilancia desproporcionada.",
    "impacto": "El malestar se sostiene fuera del horario laboral.",
}

FOCO_CLINICO = {
    "impacto": "La combinación reportada es consistente con un patrón de sobrecarga psicológica "
               "sostenida, no un evento aislado. Vale la pena explorar en sesión si existe una "
               "sensación de \"atrapamiento\" (percibir el daño con claridad, pero sin salida "
               "viable) — un patrón de riesgo conocido para sostener malestar ansioso y anímico "
               "en el tiempo.",
    "amenazas": "El predominio de amenazas y persecución sugiere explorar directamente si existe "
                "hipervigilancia — dificultad para \"desconectar\" incluso fuera del horario laboral.",
    "trato": "El predominio de trato humillante sugiere explorar el impacto en la autopercepción "
             "profesional, más allá del malestar emocional general.",
    "aislamiento": "El predominio de aislamiento sugiere explorar la red de apoyo actual — este "
                   "patrón suele acompañarse de soledad no verbalizada.",
    "presion": "El predominio de presión y sobrecarga sugiere diferenciar, en sesión, entre "
               "exigencia legítima del cargo y sobrecarga impuesta deliberadamente.",
}

PALABRAS_DEPENDENCIA = ["depend", "económic", "economic", "dinero", "desemple", "sueldo"]


def _h(doc, text, size=18):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = AZUL


def _b(doc, text, italic=False, size=11, color=CARBON):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.italic = italic


def _tabla(doc, encabezados, filas):
    t = doc.add_table(rows=1, cols=len(encabezados))
    for i, enc in enumerate(encabezados):
        t.rows[0].cells[i].text = enc
    for fila in filas:
        celdas = t.add_row().cells
        for i, valor in enumerate(fila):
            celdas[i].text = str(valor)


def _pendiente(valor, texto_si_falta):
    """Devuelve (valor mostrado, estado) -- marca 'Pendiente' si no llego el dato."""
    if valor:
        return valor, "Completo"
    return "—", f"Pendiente — {texto_si_falta}"


# ---------------------------------------------------------------- CARTA ----
def carta_tu_mirror(doc, nombre, answers):
    _h(doc, "TU MIRROR", 20)
    _b(doc, f"Hola, {nombre}.")
    texto_abierto = (answers.get("open") or "").lower()
    hay_dependencia = any(p in texto_abierto for p in PALABRAS_DEPENDENCIA)

    _b(doc, "Leímos con cuidado lo que nos contaste, y antes de entregarte los documentos "
            "técnicos, queremos decirte lo que vemos.")
    if hay_dependencia:
        _b(doc, "Lo que describes no es una sola cosa — es una combinación, y además cargas con "
                "algo que muchas veces no se dice en voz alta: el miedo a no poder sostenerte "
                "económicamente si esto se resuelve mal. Eso no es debilidad de tu parte — es "
                "exactamente el tipo de situación que hace que alguien se quede más tiempo del "
                "que debería en algo que le está haciendo daño, porque salir se siente más "
                "riesgoso que quedarse.")
        _b(doc, "Por eso este informe no termina en un solo consejo: te dejamos dos lecturas, "
                "pensadas para que no tengas que dejar tu trabajo antes de tener claridad — a "
                "menos que tú decidas que eso es lo que realmente quieres.")
    else:
        _b(doc, "Lo que describes combina varios patrones a la vez, y eso importa: no es lo "
                "mismo un hecho aislado que algo sostenido en distintos frentes. Por eso este "
                "informe no termina en un solo consejo, sino en dos lecturas — una para trabajar "
                "cómo te está afectando esto, y otra para entender qué protege la ley mientras "
                "decides con calma.")
    _b(doc, "Lo que sigue son dos documentos técnicos: uno para quien te acompañe "
            "psicológicamente, y otro para quien te oriente legalmente. Son tuyos.")
    _b(doc, "Esto lo leyó y lo armó nuestro equipo — no un sistema automático actuando solo. Si "
            "algo de lo que sigue no te hace sentido, dínoslo.", italic=True, size=10, color=GRIS)


# ---------------------------------------------------- DISCOVERY PSICOLOGICO
def discovery_psicologico(doc, nombre, answers, codigo="MRR-AUTO-A"):
    scores = compute_scores(answers)
    top, secondary = top_and_secondary(scores)

    _h(doc, f"DISCOVERY PSICOLÓGICO DE {nombre.upper()}", 16)
    _b(doc, f"Código: {codigo}", size=9, color=GRIS)

    _h(doc, "Riesgo y estado actual — lectura en 5 segundos", 13)
    riesgo = answers.get("riesgo_screening", "No recopilado")
    _b(doc, f"Riesgo agudo reportado: {riesgo}")
    _b(doc, f"Contacto de confianza dejado: {answers.get('contacto_confianza', 'No recopilado')}")
    _b(doc, f"Situación laboral actual: {answers.get('situacion_actual', 'No recopilado')}")
    _b(doc, f"En tratamiento actualmente: {answers.get('en_tratamiento', 'No recopilado')}")

    _h(doc, "Antes de la sesión — una nota para quien lo lea", 13)
    _b(doc, "La persona pagó por este proceso y puede llegar esperando que la sesión "
            "\"confirme\" lo que este documento ya describe. El rol de la primera sesión es "
            "escuchar de nuevo, no validar un resultado ya escrito.")

    _h(doc, "Motivo de consulta", 13)
    _b(doc, answers.get("motivo_consulta", "No recopilado en este envío."))

    _h(doc, "En sus propias palabras", 13)
    cita = answers.get("open", "").strip()
    if cita:
        _b(doc, f"\"{cita}\"", italic=True)

    _h(doc, "Marco técnico de referencia", 13)
    _b(doc, "Las 5 dimensiones de Mirror están inspiradas en la lógica del Cuestionario "
            "CEAL-SM/SUSESO (instrumento oficial vigente en Chile desde 2023, obligatorio bajo "
            "el Protocolo de Vigilancia de Riesgos Psicosociales del MINSAL). A modo de "
            "equivalencia aproximada:", size=10)
    _tabla(doc, ["Dimensión Mirror", "Se relaciona con (CEAL-SM/SUSESO)"], [
        ("Presión y sobrecarga", "Carga de trabajo"),
        ("Trato y comunicación / Aislamiento", "Apoyo social y calidad de liderazgo"),
        ("Impacto en ti", "Exigencias emocionales; incorpora la escala GHQ-12 de tamizaje de malestar general"),
    ])
    _b(doc, "Mirror no reemplaza el instrumento oficial, lo antecede — si se considera "
            "pertinente una medición más formal, el CEAL-SM completo está disponible "
            "públicamente a través de SUSESO.", size=9, color=GRIS)

    _h(doc, "Foco clínico preliminar (hipótesis a explorar, no una conclusión)", 13)
    _b(doc, FOCO_CLINICO[top])
    if scores["amenazas"] >= 6 and top != "amenazas":
        _b(doc, "Bandera a mirar primero: la dimensión de amenazas y persecución también "
                "puntuó alta — vale la pena explorar hipervigilancia en sesión.", size=10)
    elif scores["aislamiento"] >= 6 and top != "aislamiento":
        _b(doc, "Bandera a mirar primero: la dimensión de aislamiento también puntuó alta — "
                "vale la pena explorar la red de apoyo real en sesión.", size=10)

    _h(doc, "Resultados por dimensión", 13)
    filas = [(DIMENSION_NAMES[d], f"{level_label(scores[d])} ({scores[d]}/9)", LECTURA_DIMENSION[d])
             for d in scores]
    _tabla(doc, ["Dimensión", "Nivel", "Lectura"], filas)
    doc.add_paragraph()
    _b(doc, f"Patrón dominante: {DIMENSION_NAMES[top]}."
            + (f" Patrón secundario: {DIMENSION_NAMES[secondary]}." if secondary else ""))

    _h(doc, "Lo que falta por preguntar directamente en sesión", 13)
    faltantes = []
    if riesgo == "No recopilado":
        faltantes.append("Riesgo agudo (ideación de daño a sí mismo/a) — screening obligatorio.")
    if answers.get("situacion_actual") is None:
        faltantes.append("Situación laboral actual y si está en tratamiento hoy.")
    if answers.get("estrategias_afrontamiento") is None:
        faltantes.append("Qué ha hecho hasta ahora para manejarlo.")
    if answers.get("red_apoyo") is None:
        faltantes.append("Red de apoyo actual: ¿con quién cuenta hoy?")
    for f in faltantes:
        _b(doc, f"• {f}", size=10)

    _h(doc, "Limitaciones", 13)
    _b(doc, "Autoreporte, no evaluación clínica. El foco clínico preliminar es una hipótesis de "
            "lectura, no un diagnóstico.", size=9, color=GRIS)


# ---------------------------------------------------------- DISCOVERY LEGAL
def discovery_legal(doc, nombre, legal, codigo="MRR-AUTO-B"):
    _h(doc, f"DISCOVERY LEGAL DE {nombre.upper()}", 16)
    _b(doc, f"Código: {codigo}", size=9, color=GRIS)

    fuero = legal.get("fuero")
    finiquito = legal.get("firmo_finiquito")
    evidencia = legal.get("evidencia_lista", [])
    # Dato real: L.8 es UN campo de texto libre, no una lista de eventos ya estructurada.
    relato_libre = (legal.get("relato_libre") or "").strip()
    hay_relato = len(relato_libre) >= 80  # umbral simple: ¿hay contenido real, no solo una frase?

    urgencia = "Alta" if (fuero or (finiquito and "sí" in str(finiquito).lower())) else "Media"
    estado_proceso = legal.get("gestion_previa", "Sin gestión formal iniciada")

    _h(doc, "Triage — lectura en 5 segundos", 13)
    _b(doc, f"Urgencia: {urgencia}")
    _b(doc, "Categoría: Acoso laboral (Ley Karin)")
    _b(doc, f"Estado del proceso: {estado_proceso}")

    persona = legal.get("persona_señalada", "—")
    empresa = legal.get("empresa", "—")
    inconsistencia = persona != "—" and persona == empresa

    _h(doc, "Resumen ejecutivo", 13)
    _b(doc, f"Trabajador/a con antigüedad de {legal.get('antiguedad', '—')}, situación reportada "
            f"{legal.get('inicio', '—')} por parte de {legal.get('relacion', '—').lower() if legal.get('relacion') else '—'}. "
            f"Evidencia disponible: {', '.join(evidencia) if evidencia else 'no reportada'}.")
    if not hay_relato:
        _b(doc, "Elemento crítico pendiente: la cronología detallada de hechos no fue "
                "completada — sin ella, cualquier vía legal queda con base débil.")

    _h(doc, "Identificación de las partes", 13)
    filas = [
        ("Cargo", *_pendiente(legal.get("cargo"), "dato no recopilado")),
        ("Antigüedad", *_pendiente(legal.get("antiguedad"), "dato no recopilado")),
        ("Tipo de contrato", *_pendiente(legal.get("tipo_contrato"), "relevante para causales aplicables")),
        ("Remuneración aprox.", *_pendiente(legal.get("remuneracion"), "no recopilado")),
        ("Empresa", *_pendiente(legal.get("empresa"), "dato no recopilado")),
        ("RUT empresa", *_pendiente(legal.get("rut_empresa"), "falta RUT")),
        ("Protocolo de prevención vigente", *_pendiente(legal.get("protocolo_vigente"), "determina incumplimiento adicional")),
        ("Persona señalada", persona, "Inconsistencia: coincide con la empresa" if inconsistencia else "Completo"),
        ("Fuero (embarazo/sindical)", *_pendiente(fuero, "cambia el nivel de protección aplicable")),
        ("¿Firmó finiquito?", *_pendiente(finiquito, "crítico antes de cualquier paso siguiente")),
    ]
    _tabla(doc, ["Campo", "Dato", "Estado"], filas)

    _h(doc, "Clasificación preliminar de los hechos (hipótesis a verificar)", 13)
    _b(doc, "Con la información disponible, los hechos descritos son, en términos generales, "
            "del tipo que la Ley Karin categoriza como acoso laboral. Esta clasificación es "
            "descriptiva y preliminar — no determina si el caso califica legalmente.", size=10)

    _h(doc, "Cronología de hechos", 13)
    if hay_relato:
        _b(doc, "Relato entregado por la persona, tal como lo escribió (sin reinterpretar):")
        _b(doc, relato_libre, italic=True)
        _b(doc, f"Evidencia mencionada: {', '.join(evidencia) if evidencia else 'no reportada'}.", size=10)
    else:
        _b(doc, f"Relato entregado: \"{relato_libre or 'sin contenido'}\"", italic=True)
        _b(doc, "Esto es débil incluso bajo el estándar de indicios suficientes — se necesitan "
                "hechos concretos y ubicables en el tiempo, no una frase general. Recomendamos "
                "que la primera reunión se dedique a reconstruirla junto con el profesional.",
           size=10)

    _h(doc, "Estándar probatorio aplicable — prueba indiciaria (Art. 493 Código del Trabajo)", 13)
    _b(doc, "En tutela laboral, el trabajador no necesita probar el acoso de forma exhaustiva — "
            "basta con acreditar indicios suficientes que generen sospecha razonable. Presentados "
            "esos indicios, la carga se traslada al empleador. La Corte Suprema unificó "
            "jurisprudencia confirmando que los artículos 490 y 493 se leen de forma "
            "complementaria (rol N°12.362-2015).", size=10)
    _b(doc, "Un indicio frecuentemente subestimado: la proximidad temporal entre una queja y "
            "una consecuencia adversa (garantía de indemnidad).", size=10)

    _h(doc, "Evaluación preliminar de la evidencia disponible", 13)
    testigo_activo = legal.get("testigo_activo")  # L.9 sub-pregunta: ¿el testigo sigue en la empresa?
    catalogo = {
        "testigos": "Sí, potencialmente — un testimonio específico y creíble puede ser indicio suficiente.",
        "correos": "Preguntar explícitamente — incluso un mensaje ambiguo puede funcionar como indicio.",
        "mensajes": "Preguntar explícitamente — incluso un mensaje ambiguo puede funcionar como indicio.",
        "licencias médicas": "Pueden servir como indicio del impacto, no solo como antecedente médico.",
    }
    filas_ev = []
    for tipo, texto in catalogo.items():
        estado = "Reportado" if tipo in [e.lower() for e in evidencia] else "No reportado"
        if tipo == "testigos" and estado == "Reportado":
            if testigo_activo is None:
                estado += " — ¿sigue en la empresa? No recopilado"
            elif "no" in str(testigo_activo).lower():
                estado += " — ya no trabaja ahí (más confiable procesalmente)"
            else:
                estado += " — sigue en la empresa (puede tener miedo a represalias)"
        filas_ev.append((tipo.capitalize(), estado, texto))
    _tabla(doc, ["Elemento", "Estado", "¿Podría funcionar como indicio?"], filas_ev)

    _h(doc, "Contexto que debe orientar la estrategia", 13)
    intencion = legal.get("intencion", "")
    if "irse" in intencion.lower() or "evaluando" in intencion.lower():
        _b(doc, "La persona indica que evalúa dejar la empresa — esto abre con más naturalidad "
                "vías que terminan la relación laboral (autodespido), sin descartar tutela "
                "laboral en paralelo. A confirmar directamente en la reunión.")
    elif intencion:
        _b(doc, "La persona indica que quiere que la situación se resuelva y seguir en su "
                "puesto — esto sugiere priorizar, al menos como primer paso, vías que no "
                "requieran dejar la relación laboral (medidas de resguardo) antes que vías que "
                "sí la terminan (autodespido). A confirmar directamente con la persona, no "
                "asumir en su nombre.")
    else:
        _b(doc, "No se recopiló la intención de la persona (¿quiere seguir en el puesto, o "
                "está evaluando irse?) — es una de las primeras preguntas a hacer, porque "
                "determina qué vía legal tiene sentido priorizar.")

    _h(doc, "Vías legales a evaluar", 13)
    intencion = legal.get("intencion", "")
    if "irse" in intencion.lower() or "evaluando" in intencion.lower():
        orden = ["Autodespido (Art. 171 + Art. 160 N°1)",
                 "Tutela laboral (Art. 19 constitucional)",
                 "Medidas de resguardo dentro de la relación vigente (Art. 211-B bis)"]
    else:
        orden = ["Medidas de resguardo dentro de la relación vigente (Art. 211-B bis)",
                 "Tutela laboral (Art. 19 constitucional)",
                 "Autodespido (Art. 171 + Art. 160 N°1)"]
    for i, via in enumerate(orden, 1):
        _b(doc, f"{i}. {via}", size=10)
    _b(doc, "Plazo a verificar una vez reconstruida la cronología: 60 días hábiles desde el "
            "incumplimiento grave del empleador, si se opta por autodespido.", size=10)

    _h(doc, "Los pasos más importantes antes de avanzar", 13)
    pasos = ["Reconstruir la cronología completa con fechas."]
    if not legal.get("tipo_contrato") or not legal.get("remuneracion") or not legal.get("rut_empresa"):
        pasos.append("Confirmar tipo de contrato, remuneración aproximada y RUT de la empresa.")
    pasos.append("Confirmar identidad, disposición a declarar, y si el/los testigo(s) siguen "
                  "trabajando en la empresa.")
    if not legal.get("protocolo_vigente"):
        pasos.append("Verificar si existe protocolo de prevención vigente en la empresa (su "
                      "ausencia es un incumplimiento adicional).")
    if not fuero:
        pasos.append("Confirmar si existe fuero (embarazo, dirigente sindical) — cambia el "
                      "nivel de protección aplicable.")
    if not finiquito:
        pasos.append("Confirmar si ya firmó finiquito u otro documento de salida — crítico "
                      "antes de cualquier paso siguiente.")
    for i, paso in enumerate(pasos, 1):
        _b(doc, f"{i}. {paso}", size=10)

    _h(doc, "Limitaciones", 13)
    _b(doc, "Organiza el autoreporte y ofrece una primera lectura orientadora. No determina "
            "viabilidad legal ni reemplaza la evaluación de un abogado con el caso completo.",
       size=9, color=GRIS)


# ---------------------------------------------------------------- ANEXO DT
def anexo_dt(doc, legal):
    persona = legal.get("persona_señalada", "—")
    empresa = legal.get("empresa", "—")
    relato_libre = (legal.get("relato_libre") or "").strip()
    hay_relato = len(relato_libre) >= 80
    evidencia = legal.get("evidencia_lista", [])
    incompleto = (not hay_relato) or (persona == empresa) or (persona == "—")

    _h(doc, "ANEXO — LISTO PARA LA DIRECCIÓN DEL TRABAJO", 18)
    if incompleto:
        _b(doc, "Este anexo no presenta tu denuncia por ti, y hoy no está listo para enviar — "
                "revisa los campos marcados abajo antes de ingresar el trámite.", italic=True)
    else:
        _b(doc, "Este anexo no presenta tu denuncia por ti — te deja los campos listos para "
                "copiar en el portal Mi DT.", italic=True)

    _h(doc, "Un dato que vale la pena que sepas antes de empezar", 12)
    _b(doc, "Desde la Ley Karin, no se exige que la conducta sea repetida para que la denuncia "
            "sea admitida — un solo hecho suficientemente grave puede bastar. No subestimes tu "
            "situación por sentir que \"fue solo una vez\".", size=10)

    _h(doc, "1. Empresa", 12)
    _b(doc, f"{empresa} — RUT: {legal.get('rut_empresa', 'falta, revisar en tu contrato o liquidación')}")

    _h(doc, "2. Persona denunciada", 12)
    if persona == empresa or persona == "—":
        _b(doc, "Falta el nombre real de la persona (el dato actual coincide con la empresa o "
                "está vacío) — complétalo antes de continuar.")
    else:
        _b(doc, f"{persona} ({legal.get('relacion', '—')})")

    _h(doc, "3. Relato de los hechos — borrador redactado", 12)
    if hay_relato:
        _b(doc, relato_libre)
        _b(doc, "Este es el relato tal como lo escribiste — revísalo antes de copiarlo al "
                "formulario, y agrega fechas si te acuerdas de alguna más.", size=10)
    else:
        relato_base = (f"Desde {legal.get('relacion', 'mi jefatura').lower()}, {legal.get('inicio', 'hace un tiempo')} "
                       f"comenzó una situación que afecta mi trabajo. {relato_libre}")
        _b(doc, relato_base, italic=True)
        _b(doc, "Este borrador es solo un punto de partida — con lo que respondiste no hay "
                "suficiente detalle para redactarlo completo. Complétalo respondiendo estas tres "
                "preguntas, con tus propias palabras, y reemplaza el párrafo de arriba:", size=10)
        _b(doc, "• ¿Qué dijo o hizo exactamente, en al menos 2-3 momentos distintos?", size=10)
        _b(doc, "• ¿Hubo alguna vez testigos presentes, o alguien a quien se lo contaste después?", size=10)
        _b(doc, "• ¿Cambió algo en tu trabajo (funciones, trato, evaluaciones) después de algún "
                "reclamo o queja tuya, aunque fuera informal?", size=10)

    _h(doc, "4. Evidencia a marcar en el formulario", 12)
    opciones = ["Correos electrónicos", "Mensajes de WhatsApp u otra app", "Testigos",
                "Licencias médicas asociadas", "Evaluaciones de desempeño"]
    for op in opciones:
        marcado = "☑" if any(op.lower().split()[0] in e.lower() for e in evidencia) else "☐"
        _b(doc, f"{marcado} {op}", size=10)
    _b(doc, "Recuerda: la evidencia no es obligatoria para iniciar el trámite, pero fortalece "
            "tu denuncia — revisa si tienes algo de esto antes de marcar que no.", size=9, color=GRIS)

    _h(doc, "5. Cómo presentarla", 12)
    _b(doc, "Portal Mi DT (ClaveÚnica) > Trabajador > Denuncias y solicitudes > Denuncia "
            "Laboral y Por Vulneración de Derechos Fundamentales/Ley Karin. Orientación: "
            "600 450 4000 (lunes a viernes, 9:00 a 17:00 hrs).", size=10)

    _h(doc, "6. Qué va a pasar después de que la envíes", 12)
    _tabla(doc, ["Plazo", "Qué ocurre"], [
        ("Dentro de 48 horas", "Recibes un correo de confirmación con un código de 12 letras."),
        ("Máximo 5 días hábiles", "Debe iniciarse la investigación (interna o de la DT)."),
        ("Máximo 30 días hábiles", "Plazo total para que concluya la investigación."),
        ("15 días corridos después", "El empleador debe aplicar las medidas o sanciones que resulten."),
    ])
    _b(doc, "Guarda el código de confirmación — lo vas a necesitar para hacer seguimiento.", size=9, color=GRIS)
