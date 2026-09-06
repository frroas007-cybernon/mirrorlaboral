"""
Orquestador para la app web. Reutiliza las funciones de contenido de
generar_documentos_v3.py (carta_tu_mirror, discovery_psicologico,
discovery_legal, anexo_dt) tal cual -- no se toca ninguna palabra del
contenido. Lo unico que cambia frente al procesar_pago() original es que
en vez de guardar archivos .docx en el disco (que en un hosting serverless
no persiste), arma cada documento en memoria y devuelve los bytes, listos
para adjuntar a un correo o guardar en la base de datos.
"""
from io import BytesIO

from docx import Document

from app.mirror2_engine.generar_documentos_v3 import (
    carta_tu_mirror,
    discovery_psicologico,
    discovery_legal,
    anexo_dt,
)


def _bytes(doc: Document) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generar_documentos(tier: str, answers: dict, legal: dict) -> dict:
    """
    Devuelve un dict {clave: bytes_docx} con los documentos que corresponden
    al tier comprado:
      - "dt"       -> {"dt": ...}
      - "psico"    -> {"psicologico": ...}
      - "legal"    -> {"legal": ...}
      - "completo" -> {"completo": ..., "psicologico": ..., "legal": ..., "dt": ...}
                       (las piezas sueltas se generan además del combinado para
                       que el panel de Paz / la abogada pueda mostrar solo lo
                       que corresponde a cada una).
    """
    nombre = answers.get("nombre", "—")
    resultado = {}

    if tier == "completo":
        doc = Document()
        carta_tu_mirror(doc, nombre, answers)
        doc.add_page_break()
        discovery_psicologico(doc, nombre, answers)
        doc.add_page_break()
        discovery_legal(doc, nombre, legal)
        doc.add_page_break()
        anexo_dt(doc, legal)
        resultado["completo"] = _bytes(doc)

        doc_p = Document()
        discovery_psicologico(doc_p, nombre, answers)
        resultado["psicologico"] = _bytes(doc_p)

        doc_l = Document()
        discovery_legal(doc_l, nombre, legal)
        resultado["legal"] = _bytes(doc_l)

        doc_dt = Document()
        anexo_dt(doc_dt, legal)
        resultado["dt"] = _bytes(doc_dt)
        return resultado

    if tier == "dt":
        doc = Document()
        anexo_dt(doc, legal)
        resultado["dt"] = _bytes(doc)
    elif tier == "psico":
        doc = Document()
        discovery_psicologico(doc, nombre, answers)
        resultado["psicologico"] = _bytes(doc)
    elif tier == "legal":
        doc = Document()
        discovery_legal(doc, nombre, legal)
        resultado["legal"] = _bytes(doc)

    return resultado


NOMBRE_ARCHIVO = {
    "completo": "Mirror_Laboral_Informe_Completo.docx",
    "psicologico": "Mirror_Laboral_Discovery_Psicologico.docx",
    "legal": "Mirror_Laboral_Discovery_Legal.docx",
    "dt": "Mirror_Laboral_Anexo_DT.docx",
}
